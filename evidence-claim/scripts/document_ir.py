#!/usr/bin/env python3
"""Parse a deliverable into anchored text units.

Every unit carries an anchor precise enough to write a comment back to it
later, which is why this runs before claim extraction rather than after:
a claim that cannot be located in the source cannot be commented on.

Supported inputs: .docx, .pptx, .pdf, .md, .txt

    python3 document_ir.py report.docx
    python3 document_ir.py deck.pptx --json
    python3 document_ir.py notes.md --kinds heading,paragraph

Anchor contract, shared with write_comments.py:

    docx   {"type": "docx", "para_index": N}
           para_index counts <w:p> elements in document order inside
           word/document.xml, including paragraphs nested in tables.
    pptx   {"type": "pptx", "slide_index": N, "shape_id": N, "para_index": N}
    pdf    {"type": "pdf", "page": N, "y": float|null, "line_index": N}
    text   {"type": "text", "line_start": N, "line_end": N}

Comments anchor to the containing paragraph, not to a character span. A
sub-paragraph span would need run splitting in Word, and the precision buys
nothing: a comment on the paragraph holding the claim reads the same.
"""

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

TEXT_EXTS = {".md", ".txt", ".markdown", ".mdown"}


def _need(module, ext):
    try:
        return __import__(module)
    except ImportError:
        sys.exit(
            f"error: reading {ext} needs the '{module}' package "
            f"(pip3 install {module})"
        )


def _lxml():
    """lxml rather than the standard library, deliberately.

    ElementTree re-derives namespace prefixes when it writes, which can rename
    the prefixes a real Word file declares. Word's mc:Ignorable attribute names
    those prefixes as a string, so a rename leaves it pointing at prefixes that
    no longer exist and Word reads the file as corrupt. lxml round-trips them.
    """
    try:
        from lxml import etree

        return etree
    except ImportError:
        sys.exit(
            "error: lxml is required for .docx handling.\n"
            "       pip3 install lxml"
        )


def _parser(etree):
    """Untrusted Office XML: no entity expansion, no network, no DTD loading."""
    return etree.XMLParser(
        resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False
    )


# --------------------------------------------------------------------------
# docx
# --------------------------------------------------------------------------


def parse_docx(path):
    etree = _lxml()

    with zipfile.ZipFile(path) as z:
        try:
            xml = z.read("word/document.xml")
        except KeyError:
            sys.exit(f"error: {path} has no word/document.xml; is it a real .docx?")
    root = etree.fromstring(xml, _parser(etree))

    units = []
    section = None
    for idx, p in enumerate(root.iter(f"{W}p")):
        text = "".join(t.text or "" for t in p.iter(f"{W}t")).strip()
        if not text:
            continue

        style = ""
        pstyle = p.find(f"{W}pPr/{W}pStyle")
        if pstyle is not None:
            style = pstyle.get(f"{W}val", "") or ""

        in_table = any(
            a.tag == f"{W}tc" for a in p.iterancestors()
        )
        numbered = p.find(f"{W}pPr/{W}numPr") is not None

        if style.startswith("Heading") or style == "Title":
            kind = "heading"
            section = text
        elif in_table:
            kind = "table_cell"
        elif numbered or style.startswith("List"):
            kind = "list_item"
        else:
            kind = "paragraph"

        units.append(
            {
                "unit_id": f"d{idx}",
                "kind": kind,
                "text": text,
                "order": len(units),
                "anchor": {"type": "docx", "para_index": idx},
                "context": {"section": section},
            }
        )
    return units


# --------------------------------------------------------------------------
# pptx
# --------------------------------------------------------------------------


def parse_pptx(path):
    _need("pptx", ".pptx")
    from pptx import Presentation

    prs = Presentation(str(path))
    units = []
    for s_idx, slide in enumerate(prs.slides):
        title_shape = None
        try:
            title_shape = slide.shapes.title
        except (AttributeError, ValueError):
            title_shape = None
        title_text = (title_shape.text.strip() if title_shape is not None else "") or None

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            is_title = title_shape is not None and shape.shape_id == title_shape.shape_id
            for p_idx, para in enumerate(shape.text_frame.paragraphs):
                text = "".join(r.text for r in para.runs).strip()
                if not text:
                    continue
                units.append(
                    {
                        "unit_id": f"s{s_idx}_{shape.shape_id}_{p_idx}",
                        "kind": "slide_title" if is_title else "slide_body",
                        "text": text,
                        "order": len(units),
                        "anchor": {
                            "type": "pptx",
                            "slide_index": s_idx,
                            "shape_id": shape.shape_id,
                            "para_index": p_idx,
                        },
                        "context": {"slide": s_idx + 1, "slide_title": title_text},
                    }
                )

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            for p_idx, para in enumerate(notes.paragraphs):
                text = "".join(r.text for r in para.runs).strip()
                if not text:
                    continue
                units.append(
                    {
                        "unit_id": f"s{s_idx}_notes_{p_idx}",
                        "kind": "notes",
                        "text": text,
                        "order": len(units),
                        "anchor": {
                            "type": "pptx",
                            "slide_index": s_idx,
                            "shape_id": None,
                            "para_index": p_idx,
                            "notes": True,
                        },
                        "context": {"slide": s_idx + 1, "slide_title": title_text},
                    }
                )
    return units


# --------------------------------------------------------------------------
# pdf
# --------------------------------------------------------------------------


def parse_pdf(path):
    _need("pypdf", ".pdf")
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    units = []
    for page_no, page in enumerate(reader.pages, start=1):
        lines = []

        def visitor(text, cm, tm, font_dict, font_size, _lines=lines):
            if not text.strip():
                return
            _lines.append((round(float(tm[5]), 1), text))

        try:
            page.extract_text(visitor_text=visitor)
        except Exception:
            lines = []

        if lines:
            grouped = {}
            for y, text in lines:
                grouped.setdefault(y, []).append(text)
            ordered = sorted(grouped.items(), key=lambda kv: -kv[0])
            blocks = [(y, "".join(parts).strip()) for y, parts in ordered]
        else:
            raw = page.extract_text() or ""
            blocks = [(None, ln.strip()) for ln in raw.splitlines()]

        for line_index, (y, text) in enumerate(blocks):
            if not text:
                continue
            units.append(
                {
                    "unit_id": f"p{page_no}_{line_index}",
                    "kind": "page_line",
                    "text": text,
                    "order": len(units),
                    "anchor": {
                        "type": "pdf",
                        "page": page_no,
                        "y": y,
                        "line_index": line_index,
                    },
                    "context": {"page": page_no},
                }
            )
    return units


# --------------------------------------------------------------------------
# markdown / plain text
# --------------------------------------------------------------------------


HEADING_RE = re.compile(r"^\s{0,3}(#{1,6})\s+(.*\S)\s*$")
BULLET_RE = re.compile(r"^\s*([-*+]|\d+[.)])\s+\S")
BULLET_STRIP_RE = re.compile(r"^\s*([-*+]|\d+[.)])\s+")


def parse_text(path):
    raw = Path(path).read_text(encoding="utf-8", errors="replace")
    lines = raw.splitlines()
    units = []
    section = None

    block = []
    block_start = 0

    def flush(end_line):
        nonlocal block, block_start
        if not block:
            return
        text = " ".join(s.strip() for s in block).strip()
        if text:
            is_bullet = bool(BULLET_RE.match(block[0]))
            kind = "list_item" if is_bullet else "paragraph"
            if is_bullet:
                text = BULLET_STRIP_RE.sub("", text, count=1).strip() or text
            units.append(
                {
                    "unit_id": f"t{block_start + 1}",
                    "kind": kind,
                    "text": text,
                    "order": len(units),
                    "anchor": {
                        "type": "text",
                        "line_start": block_start + 1,
                        "line_end": end_line,
                    },
                    "context": {"section": section},
                }
            )
        block = []

    for i, line in enumerate(lines):
        h = HEADING_RE.match(line)
        if h:
            flush(i)
            section = h.group(2)
            units.append(
                {
                    "unit_id": f"t{i + 1}",
                    "kind": "heading",
                    "text": section,
                    "order": len(units),
                    "anchor": {"type": "text", "line_start": i + 1, "line_end": i + 1},
                    "context": {"section": section, "level": len(h.group(1))},
                }
            )
            continue
        if not line.strip():
            flush(i)
            continue
        if BULLET_RE.match(line):
            # Each bullet is its own claim candidate, so it gets its own unit
            # rather than being merged into the surrounding block.
            flush(i)
        if not block:
            block_start = i
        block.append(line)
    flush(len(lines))
    return units


# --------------------------------------------------------------------------


PARSERS = {".docx": parse_docx, ".pptx": parse_pptx, ".pdf": parse_pdf}

# Formats that accept a real comment layer written back into the file.
COMMENTABLE = {".docx", ".pptx", ".pdf"}


def parse(path):
    path = Path(path)
    if not path.exists():
        sys.exit(f"error: no such file: {path}")
    ext = path.suffix.lower()
    if ext in PARSERS:
        return PARSERS[ext](path)
    if ext in TEXT_EXTS:
        return parse_text(path)
    sys.exit(
        f"error: unsupported deliverable type '{ext}'. "
        f"Supported: {', '.join(sorted(set(PARSERS) | TEXT_EXTS))}"
    )


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("path", help="deliverable to parse")
    ap.add_argument("--json", action="store_true", help="emit the full unit list as JSON")
    ap.add_argument(
        "--kinds",
        help="comma-separated unit kinds to keep (default: all)",
    )
    ap.add_argument(
        "--min-words",
        type=int,
        default=0,
        help="drop units shorter than this many words",
    )
    args = ap.parse_args()

    units = parse(args.path)
    if args.kinds:
        keep = {k.strip() for k in args.kinds.split(",") if k.strip()}
        units = [u for u in units if u["kind"] in keep]
    if args.min_words:
        units = [u for u in units if len(u["text"].split()) >= args.min_words]

    ext = Path(args.path).suffix.lower()
    if args.json:
        print(
            json.dumps(
                {
                    "file": str(args.path),
                    "format": ext.lstrip("."),
                    "commentable": ext in COMMENTABLE,
                    "unit_count": len(units),
                    "units": units,
                },
                indent=2,
            )
        )
        return

    print(f"{args.path}")
    print(f"  format:      {ext.lstrip('.')}")
    print(f"  commentable: {'yes' if ext in COMMENTABLE else 'no (report only)'}")
    print(f"  units:       {len(units)}")
    counts = {}
    for u in units:
        counts[u["kind"]] = counts.get(u["kind"], 0) + 1
    for kind, n in sorted(counts.items(), key=lambda kv: -kv[1]):
        print(f"    {kind:<12} {n}")
    print()
    for u in units[:12]:
        preview = u["text"][:88] + ("..." if len(u["text"]) > 88 else "")
        print(f"  [{u['unit_id']}] {u['kind']}: {preview}")
    if len(units) > 12:
        print(f"  ... {len(units) - 12} more")


if __name__ == "__main__":
    main()
