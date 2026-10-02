#!/usr/bin/env python3
"""Build-gating assertions for evidence-check.

Exercises the three scripts end to end against files built here, so a broken
parser, a broken registry, or a comment writer that produces an unopenable
document fails the package build rather than a real audit.

Run directly: python3 fixtures/test_evidence_check.py
"""

import json
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"
DOC_IR = SCRIPTS / "document_ir.py"
INVENTORY = SCRIPTS / "corpus_inventory.py"
COMMENTS = SCRIPTS / "write_comments.py"

FAILURES = []
SKIPS = []


def check(label, condition, detail=""):
    if condition:
        print(f"  ok    {label}")
    else:
        print(f"  FAIL  {label} {detail}")
        FAILURES.append(label)


def run(*args, expect=0):
    proc = subprocess.run(
        [sys.executable, *[str(a) for a in args]],
        capture_output=True,
        text=True,
    )
    if expect is not None and proc.returncode != expect:
        print(proc.stdout)
        print(proc.stderr, file=sys.stderr)
        raise AssertionError(
            f"{args[0]} exited {proc.returncode}, expected {expect}"
        )
    return proc


CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""

ROOT_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def make_docx(path):
    """A minimal but real .docx, built without python-docx so the fixture has
    no dependency the scripts themselves do not already need."""

    def para(text, style=None):
        ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
        return f"<w:p>{ppr}<w:r><w:t>{text}</w:t></w:r></w:p>"

    body = (
        para("Findings", "Heading1")
        + para("Most operators work around the scheduling tool.")
        + para("We interviewed 11 operators across 4 sites.")
        + "<w:tbl><w:tr><w:tc>"
        + para("Cited by 3 of 11")
        + "</w:tc></w:tr></w:tbl>"
        + para("Rebuild the handoff, because schedulers cannot see crew status.")
    )
    document = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{W_NS}"><w:body>{body}</w:body></w:document>'
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", CONTENT_TYPES)
        z.writestr("_rels/.rels", ROOT_RELS)
        z.writestr("word/document.xml", document)


def test_document_ir(tmp):
    print("document_ir")
    docx = tmp / "report.docx"
    make_docx(docx)
    out = run(DOC_IR, docx, "--json").stdout
    data = json.loads(out)
    kinds = [u["kind"] for u in data["units"]]
    texts = [u["text"] for u in data["units"]]

    check("docx parses every non-empty paragraph", data["unit_count"] == 5, kinds)
    check("docx heading detected", kinds[0] == "heading")
    check("docx table cell detected", "table_cell" in kinds)
    check("docx reports commentable", data["commentable"] is True)
    check(
        "docx anchors are paragraph indexes",
        all(u["anchor"]["type"] == "docx" for u in data["units"])
        and [u["anchor"]["para_index"] for u in data["units"]] == [0, 1, 2, 3, 4],
        [u["anchor"] for u in data["units"]],
    )
    check("docx heading carries into section context", data["units"][1]["context"]["section"] == "Findings")

    md = tmp / "notes.md"
    md.write_text(
        "# Synthesis\n\nMost operators work around the tool.\n\n"
        "- 3 of 11 said so\n- no participant mentioned training\n",
        encoding="utf-8",
    )
    data = json.loads(run(DOC_IR, md, "--json").stdout)
    bullets = [u for u in data["units"] if u["kind"] == "list_item"]
    check("markdown splits one unit per bullet", len(bullets) == 2, [u["text"] for u in bullets])
    check("markdown strips the bullet marker", bullets[0]["text"] == "3 of 11 said so", bullets[0]["text"])
    check("markdown reports not commentable", data["commentable"] is False)

    unsupported = tmp / "thing.xyz"
    unsupported.write_text("x", encoding="utf-8")
    proc = run(DOC_IR, unsupported, expect=1)
    check("unsupported format exits 1 with a reason", "unsupported deliverable type" in proc.stderr)


def test_entity_hardening(tmp):
    """A DOCTYPE with an entity must not be expanded into the parsed text."""
    evil = tmp / "evil.docx"
    make_docx(evil)
    with zipfile.ZipFile(evil) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    doc = parts["word/document.xml"].decode("utf-8")
    doc = doc.replace(
        "<w:document",
        '<!DOCTYPE d [<!ENTITY boom "EXPANDED-ENTITY-TEXT">]><w:document', 1
    ).replace("Findings", "&boom;", 1)
    parts["word/document.xml"] = doc.encode("utf-8")
    with zipfile.ZipFile(evil, "w") as z:
        for n, b in parts.items():
            z.writestr(n, b)
    proc = run(DOC_IR, evil, "--json", expect=None)
    check("entity in .docx is not expanded", "EXPANDED-ENTITY-TEXT" not in proc.stdout, proc.stdout[:200])


def test_inventory(tmp):
    print("corpus_inventory")
    proj = tmp / "proj"
    (proj / "transcripts").mkdir(parents=True)
    (proj / "drafts").mkdir()
    (proj / "transcripts" / "p04-interview.md").write_text("I just message the supervisor", encoding="utf-8")
    (proj / "transcripts" / "p07.vtt").write_text("WEBVTT", encoding="utf-8")
    (proj / "drafts" / "old.md").write_text("draft", encoding="utf-8")
    make_docx(proj / "report.docx")

    run(INVENTORY, "init", proj, "--name", "fixture", "--unit", "participants")
    check("registry created", (proj / ".evidence-check" / "registry.json").exists())

    run(INVENTORY, "scan", proj, "--deliverable", "report.docx")
    data = json.loads(run(INVENTORY, "manifest", proj, "--json").stdout)
    locators = {s["locator"] for s in data["present"]}
    check("transcripts registered", "transcripts/p04-interview.md" in locators)
    check("vtt classified as transcript", any(s["kind"] == "transcript" for s in data["present"] if s["locator"].endswith(".vtt")))
    check("deliverable excluded from its own evidence", "report.docx" not in locators, locators)
    check("registry itself excluded", not any(l.startswith(".evidence-check") for l in locators))
    check("file sources carry a resolvable link", all(s["link"].startswith("file://") for s in data["present"]))

    run(INVENTORY, "exclude", proj, "--pattern", "drafts/**", "--reason", "working drafts")
    proc = run(INVENTORY, "exclude", proj, "--pattern", "drafts/**", "--reason", "again", expect=1)
    check("duplicate exclusion refused", "already excluded" in proc.stderr)

    run(INVENTORY, "add", proj, "--type", "miro", "--locator", "board123",
        "--label", "Synthesis board", "--link", "https://miro.com/app/board/board123/", "--kind", "board")
    data = json.loads(run(INVENTORY, "manifest", proj, "--json").stdout)
    check("connector source registered", any(s["type"] == "miro" for s in data["present"]))
    check("excluded drafts gone from manifest", not any("drafts/" in s["locator"] for s in data["present"]))

    (proj / "transcripts" / "p09-interview.md").write_text("nobody touches it", encoding="utf-8")
    (proj / "transcripts" / "p07.vtt").unlink()
    out = run(INVENTORY, "scan", proj, "--deliverable", "report.docx").stdout
    check("new file reported as new", "p09-interview.md" in out and "new since last run" in out)
    check("removed file reported as missing", "missing since last run" in out and "p07.vtt" in out)
    check("exclusion list printed every run", "working drafts" in out)
    check("unit of analysis printed every run", "participants" in out)

    data = json.loads(run(INVENTORY, "manifest", proj, "--json").stdout)
    check("unit survives scan and add", data["unit_of_analysis"] == "participants", data["unit_of_analysis"])

    run(INVENTORY, "set-unit", proj, "--unit", "reports")
    data = json.loads(run(INVENTORY, "manifest", proj, "--json").stdout)
    check("set-unit changes the denominator", data["unit_of_analysis"] == "reports", data["unit_of_analysis"])

    # A registry written before unit_of_analysis existed must still read.
    legacy = tmp / "legacy"
    (legacy / ".evidence-check").mkdir(parents=True)
    reg = json.loads((proj / ".evidence-check" / "registry.json").read_text(encoding="utf-8"))
    reg.pop("unit_of_analysis")
    reg["root"] = str(legacy)
    (legacy / ".evidence-check" / "registry.json").write_text(json.dumps(reg), encoding="utf-8")
    data = json.loads(run(INVENTORY, "manifest", legacy, "--json").stdout)
    check("registry without the field falls back", data["unit_of_analysis"] == "sources", data["unit_of_analysis"])


def test_write_comments(tmp):
    print("write_comments")
    docx = tmp / "target.docx"
    make_docx(docx)
    verdicts = tmp / "verdicts.json"
    verdicts.write_text(
        json.dumps(
            [
                {"claim_id": "C01", "anchor": {"type": "docx", "para_index": 1},
                 "comment": "overstated\nsupport 3 - disconfirm 2 - silent 6"},
                {"claim_id": "C02", "anchor": {"type": "docx", "para_index": 4},
                 "comment": "unsupported\nnothing in the corpus bears on this"},
            ]
        ),
        encoding="utf-8",
    )

    out = run(COMMENTS, docx, verdicts, "--dry-run").stdout
    check("dry run writes nothing", "would be written" in out and "word/comments.xml" not in zipfile.ZipFile(docx).namelist())

    checked = tmp / "target-checked.docx"
    run(COMMENTS, docx, verdicts, "-o", checked)
    z = zipfile.ZipFile(checked)
    comments = z.read("word/comments.xml").decode()
    document = z.read("word/document.xml").decode()
    ctypes = z.read("[Content_Types].xml").decode()
    rels = z.read("word/_rels/document.xml.rels").decode()

    check("two comments written", comments.count("<w:comment ") == 2)
    check("comment text preserved", "support 3 - disconfirm 2 - silent 6" in comments)
    check("range markers anchor the comments", document.count("commentRangeStart") == 2 and document.count("commentRangeEnd") == 2)
    check("comment reference run present", document.count("commentReference") == 2)
    check("content type override added", "/word/comments.xml" in ctypes)
    check("relationship added", "comments.xml" in rels)

    run(COMMENTS, checked, verdicts, "--in-place")
    again = zipfile.ZipFile(checked).read("word/comments.xml").decode()
    check("re-run appends and preserves existing comments", again.count("<w:comment ") == 4, again.count("<w:comment "))
    ids = [line.split('w:id="')[1].split('"')[0] for line in again.split("<w:comment ")[1:]]
    check("comment ids stay unique", len(set(ids)) == 4, ids)

    bad = tmp / "bad.json"
    bad.write_text(json.dumps([{"claim_id": "C99", "anchor": {"type": "docx", "para_index": 999}, "comment": "x"}]), encoding="utf-8")
    proc = run(COMMENTS, docx, bad, "-o", tmp / "bad.docx", expect=1)
    check("out-of-range anchor is skipped and reported", "outside 0.." in proc.stdout)

    malformed = tmp / "malformed.json"
    malformed.write_text(json.dumps([{"claim_id": "C1"}]), encoding="utf-8")
    proc = run(COMMENTS, docx, malformed, expect=1)
    check("verdict missing a field is refused", "missing 'anchor'" in proc.stderr)

    md = tmp / "notes.md"
    md.write_text("# x\n\nclaim\n", encoding="utf-8")
    proc = run(COMMENTS, md, verdicts, expect=1)
    check("format with no comment layer is refused", "no comment layer" in proc.stderr)


def test_real_word_file(tmp):
    """lxml is a dependency specifically to keep this true.

    A Word-written document.xml declares ~17 namespace prefixes and names some
    of them in mc:Ignorable. A writer that re-derives prefixes can rename them,
    leaving mc:Ignorable pointing at prefixes that no longer exist, which Word
    reads as corruption. The handcrafted docx elsewhere in this fixture has one
    namespace and cannot catch that, so this test uses a real one.
    """
    print("real Word file round trip")
    try:
        from docx import Document
    except ImportError:
        SKIPS.append("real Word round trip (python-docx not installed)")
        print("  skip  python-docx not installed")
        return

    import re

    src = tmp / "real.docx"
    d = Document()
    d.add_heading("Findings", 1)
    d.add_paragraph("Most operators work around the scheduling tool.")
    d.save(str(src))

    v = tmp / "vreal.json"
    v.write_text(
        json.dumps([{"claim_id": "C01", "anchor": {"type": "docx", "para_index": 1},
                     "comment": "overstated"}]),
        encoding="utf-8",
    )
    out = tmp / "real-checked.docx"
    run(COMMENTS, src, v, "-o", out)

    def root_attrs(path):
        x = zipfile.ZipFile(path).read("word/document.xml").decode()
        i = x.index("<w:document")
        head = x[i:x.index(">", i) + 1]
        ignorable = re.search(r'mc:Ignorable="([^"]*)"', head)
        return sorted(re.findall(r"xmlns:(\w+)=", head)), (
            ignorable.group(1).split() if ignorable else []
        ), x

    before, ig_before, _ = root_attrs(src)
    after, ig_after, body = root_attrs(out)

    check("namespace prefixes preserved exactly", before == after, f"{before} != {after}")
    check("the file really does declare many prefixes", len(after) > 5, after)
    check("mc:Ignorable preserved", ig_before == ig_after, f"{ig_before} != {ig_after}")
    check(
        "every mc:Ignorable prefix is still declared",
        all(p in after for p in ig_after),
        [p for p in ig_after if p not in after],
    )
    check("no prefix was auto-renamed", "ns0:" not in body)


def test_optional_formats(tmp):
    print("optional formats")
    try:
        from pptx import Presentation
    except ImportError:
        SKIPS.append("pptx (python-pptx not installed)")
        print("  skip  pptx, python-pptx not installed")
    else:
        deck = tmp / "deck.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Most techs bypass the scheduler"
        slide.placeholders[1].text_frame.text = "3 of 11 described a workaround"
        prs.save(str(deck))

        data = json.loads(run(DOC_IR, deck, "--json").stdout)
        kinds = [u["kind"] for u in data["units"]]
        check("pptx title extracted as a claim unit", "slide_title" in kinds, kinds)

        v = tmp / "vp.json"
        v.write_text(json.dumps([{"claim_id": "C01", "anchor": {"type": "pptx", "slide_index": 0}, "comment": "overstated"}]), encoding="utf-8")
        outdeck = tmp / "deck-checked.pptx"
        run(COMMENTS, deck, v, "-o", outdeck)
        names = zipfile.ZipFile(outdeck).namelist()
        check("pptx comment part written", any("comments/comment" in n for n in names), names[:5])
        check("pptx author list written", "ppt/commentAuthors.xml" in names)
        Presentation(str(outdeck))
        check("pptx reopens after annotation", True)
        proc = run(COMMENTS, deck, v, "--in-place", expect=1)
        check("pptx in-place refused", "--in-place is refused" in proc.stderr)

    try:
        from pypdf import PdfReader, PdfWriter
        from pypdf.annotations import Text  # noqa: F401  needs pypdf 3+
    except ImportError:
        SKIPS.append("pdf (pypdf 3+ with annotations not installed)")
        print("  skip  pdf, pypdf 3+ with annotation support not installed")
        return

    pdf = tmp / "doc.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with open(pdf, "wb") as fh:
        writer.write(fh)

    v = tmp / "vpdf.json"
    v.write_text(json.dumps([{"claim_id": "C01", "anchor": {"type": "pdf", "page": 1, "y": 600.0}, "comment": "overstated"}]), encoding="utf-8")
    outpdf = tmp / "doc-checked.pdf"
    run(COMMENTS, pdf, v, "-o", outpdf)
    annots = PdfReader(str(outpdf)).pages[0].get("/Annots")
    check("pdf annotation attached to the page", bool(annots) and len(annots) == 1)
    check("pdf annotation carries the verdict text", "overstated" in str(annots[0].get_object().get("/Contents")))


def main():
    for script in (DOC_IR, INVENTORY, COMMENTS):
        if not script.exists():
            sys.exit(f"missing script: {script}")

    try:
        import lxml  # noqa: F401
    except ImportError:
        sys.exit(
            "error: lxml is required by evidence-check and is not installed for\n"
            f"       {sys.executable}\n"
            "       pip3 install lxml"
        )

    tmp = Path(tempfile.mkdtemp(prefix="evidence-check-fixture-"))
    try:
        test_document_ir(tmp)
        test_entity_hardening(tmp)
        test_inventory(tmp)
        test_write_comments(tmp)
        test_real_word_file(tmp)
        test_optional_formats(tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    print()
    if SKIPS:
        print(f"skipped: {', '.join(SKIPS)}")
    if FAILURES:
        print(f"FAILED {len(FAILURES)}: {', '.join(FAILURES)}")
        sys.exit(1)
    print("all assertions passed")


if __name__ == "__main__":
    main()
